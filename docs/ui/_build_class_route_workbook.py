#!/usr/bin/env python3
"""Simple Class-route structure workbook: bilingual text for design."""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = "/workspace/docs/ui/设计结构呈现.xlsx"

INK = "1F2937"
WHITE = "FFFFFF"
NAVY = "1B2A4A"
LINE = "D9DEE8"
PEACH = "FFF6EE"
GOLD = "FFF6D6"
BLUE = "E8F2FC"
GREY = "F4F6FA"

LEVEL_FILL = {
    "Level 1": "5B5BD6",
    "Level 2": "2878D0",
    "Level 3": "0E9F6E",
    "Level 4": "E8791A",
    "Level 5": "C9A227",
    "Level 6": "7C3AED",
}

TYPE_EN = {
    "词汇课": "Vocabulary",
    "句型课": "Sentence Pattern",
    "主题阅读课": "Themed Reading",
    "字母课": "Alphabet",
    "拼读阅读课": "Phonics Reading",
    "复习测评课": "Review & Assessment",
    "自拼课": "Independent Phonics",
    "语法阅读课": "Grammar Reading",
    "词句听说课": "Words & Sentences (Listen & Speak)",
    "阅读基础课": "Reading Foundations",
    "语法课": "Grammar",
    "篇章阅读课": "Passage Reading",
    "写作课": "Writing",
    "语言筑基课": "Language Foundations",
}

# Short labels for the route chip. 1 word EN / 2 characters ZH when possible.
ROUTE_ZH = {
    "词汇课": "单词",
    "句型课": "句子",
    "主题阅读课": "阅读",
    "字母课": "字母",
    "拼读阅读课": "拼读",
    "复习测评课": "测评",
    "自拼课": "自拼",
    "语法阅读课": "语法读",
    "词句听说课": "听说",
    "阅读基础课": "基础",
    "语法课": "语法",
    "篇章阅读课": "篇章",
    "写作课": "写作",
    "语言筑基课": "筑基",
}

ROUTE_EN = {
    "词汇课": "Words",
    "句型课": "Sentences",
    "主题阅读课": "Story",
    "字母课": "ABC",
    "拼读阅读课": "Phonics",
    "复习测评课": "Quiz",
    "自拼课": "Blend",
    "语法阅读课": "Read",
    "词句听说课": "Speak",
    "阅读基础课": "Basics",
    "语法课": "Grammar",
    "篇章阅读课": "Passage",
    "写作课": "Writing",
    "语言筑基课": "Core",
}

CATALOG = [
    ("Level 1", "Lesson 1", "词汇课"),
    ("Level 1", "Lesson 2", "句型课"),
    ("Level 1", "Lesson 3", "主题阅读课"),
    ("Level 1", "Lesson 4", "字母课"),
    ("Level 1", "Lesson 5", "拼读阅读课"),
    ("Level 1", "Lesson 6", "复习测评课"),
    ("Level 2", "Lesson 1", "词汇课"),
    ("Level 2", "Lesson 2", "句型课"),
    ("Level 2", "Lesson 3", "自拼课"),
    ("Level 2", "Lesson 4", "拼读阅读课"),
    ("Level 2", "Lesson 5", "语法阅读课"),
    ("Level 2", "Lesson 6", "复习测评课"),
    ("Level 3", "Lesson 1", "词句听说课"),
    ("Level 3", "Lesson 2", "词句听说课"),
    ("Level 3", "Lesson 3", "词句听说课"),
    ("Level 3", "Lesson 4", "阅读基础课"),
    ("Level 3", "Lesson 5", "语法课"),
    ("Level 3", "Lesson 6", "复习测评课"),
    ("Level 4", "Lesson 1", "词句听说课"),
    ("Level 4", "Lesson 2", "篇章阅读课"),
    ("Level 4", "Lesson 3", "词句听说课"),
    ("Level 4", "Lesson 4", "篇章阅读课"),
    ("Level 4", "Lesson 5", "阅读基础课"),
    ("Level 4", "Lesson 6", "复习测评课"),
    ("Level 4", "Lesson 7", "词句听说课"),
    ("Level 4", "Lesson 8", "篇章阅读课"),
    ("Level 4", "Lesson 9", "词句听说课"),
    ("Level 4", "Lesson 10", "篇章阅读课"),
    ("Level 4", "Lesson 11", "写作课"),
    ("Level 4", "Lesson 12", "复习测评课"),
    ("Level 5", "U1L1", "词句听说课"),
    ("Level 5", "U1L2", "篇章阅读课"),
    ("Level 5", "U1L3", "词句听说课"),
    ("Level 5", "U1L4", "篇章阅读课"),
    ("Level 5", "U1L5", "语言筑基课"),
    ("Level 5", "U1L6", "复习测评课"),
    ("Level 5", "U1L7", "词句听说课"),
    ("Level 5", "U1L8", "篇章阅读课"),
    ("Level 5", "U1L9", "词句听说课"),
    ("Level 5", "U1L10", "篇章阅读课"),
    ("Level 5", "U1L11", "写作课"),
    ("Level 5", "U1L12", "复习测评课"),
    ("Level 6", "U1L1", "词句听说课"),
    ("Level 6", "U1L2", "篇章阅读课"),
    ("Level 6", "U1L3", "词句听说课"),
    ("Level 6", "U1L4", "篇章阅读课"),
    ("Level 6", "U1L5", "语言筑基课"),
    ("Level 6", "U1L6", "复习测评课"),
    ("Level 6", "U1L7", "词句听说课"),
    ("Level 6", "U1L8", "篇章阅读课"),
    ("Level 6", "U1L9", "词句听说课"),
    ("Level 6", "U1L10", "篇章阅读课"),
    ("Level 6", "U1L11", "写作课"),
    ("Level 6", "U1L12", "复习测评课"),
]

thin = Border(
    left=Side(style="thin", color=LINE),
    right=Side(style="thin", color=LINE),
    top=Side(style="thin", color=LINE),
    bottom=Side(style="thin", color=LINE),
)


def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)


def style(cell, *, bg=WHITE, fg=INK, bold=False, size=12):
    cell.fill = fill(bg)
    cell.font = Font(name="Calibri", size=size, bold=bold, color=fg)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = thin


def header_row(ws, headers):
    for i, text in enumerate(headers, 1):
        cell = ws.cell(1, i, text)
        style(cell, bg=NAVY, fg=WHITE, bold=True, size=12)
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False


def main():
    wb = Workbook()

    # --- Route ---
    ws = wb.active
    ws.title = "Route 路线"
    header_row(ws, ["Phase", "阶段", "Step", "步骤"])
    for col, width in enumerate([18, 14, 18, 14], 1):
        ws.column_dimensions[get_column_letter(col)].width = width

    steps = [
        ("In class", "课中", "Preview", "预习", PEACH),
        ("In class", "课中", "Teaching", "教学", PEACH),
        ("In class", "课中", "Practice", "练习", PEACH),
        ("In class", "课中", "Summary", "总结", PEACH),
        ("In class", "课中", "Award", "领奖", GOLD),
        ("After class", "课后", "Review", "复习", BLUE),
        ("After class", "课后", "Speaking", "口语", BLUE),
        ("After class", "课后", "Listening", "听力", BLUE),
        ("After class", "课后", "Report", "报告", BLUE),
    ]
    for i, (pe, pz, se, sz, bg) in enumerate(steps, 2):
        for col, val in enumerate((pe, pz, se, sz), 1):
            style(ws.cell(i, col, val), bg=bg, bold=(col >= 3), size=13)
        ws.row_dimensions[i].height = 32

    # --- Catalog: short names on the route ---
    cat = wb.create_sheet("Lesson catalog 课表")
    header_row(
        cat,
        ["Level 级别", "Lesson 课次", "课型", "Full name EN", "On route 中文", "On route EN"],
    )
    for col, width in enumerate([14, 14, 16, 34, 16, 14], 1):
        cat.column_dimensions[get_column_letter(col)].width = width

    for i, (level, lesson, typ) in enumerate(CATALOG, 2):
        bg = LEVEL_FILL[level]
        fg = NAVY if level == "Level 5" else WHITE
        stripe = WHITE if i % 2 == 0 else GREY
        vals = [level, lesson, typ, TYPE_EN[typ], ROUTE_ZH[typ], ROUTE_EN[typ]]
        for col, val in enumerate(vals, 1):
            cell = cat.cell(i, col, val)
            if col <= 2:
                style(cell, bg=bg, fg=fg, bold=True)
            elif col >= 5:
                style(cell, bg=PEACH, bold=True)
            else:
                style(cell, bg=stripe)
        cat.row_dimensions[i].height = 22
    cat.auto_filter.ref = f"A1:F{1 + len(CATALOG)}"

    # --- Unique short labels for design ---
    tags = wb.create_sheet("Type tags 课型标签")
    header_row(
        tags,
        ["课型", "Full name EN", "On route 中文", "On route EN"],
    )
    for col, width in enumerate([16, 36, 16, 14], 1):
        tags.column_dimensions[get_column_letter(col)].width = width
    seen = []
    for typ in TYPE_EN:
        seen.append(typ)
    for i, typ in enumerate(seen, 2):
        stripe = WHITE if i % 2 == 0 else GREY
        for col, val in enumerate(
            (typ, TYPE_EN[typ], ROUTE_ZH[typ], ROUTE_EN[typ]), 1
        ):
            style(tags.cell(i, col, val), bg=PEACH if col >= 3 else stripe, bold=(col >= 3))
        tags.row_dimensions[i].height = 26
    tags.auto_filter.ref = f"A1:D{1 + len(seen)}"

    wb.properties.title = "Class route structure 上课路线结构"
    wb.save(OUT)
    print("Wrote", OUT)


if __name__ == "__main__":
    main()
